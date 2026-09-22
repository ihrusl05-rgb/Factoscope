"""Генерация Excel-шаблона для копирайтера.

Пример:
    python manage.py make_template --month 2026-10
    → horoscope_2026-10.xlsx (даты месяца x 12 знаков)
"""

import os
from calendar import monthrange
from datetime import date

from django.core.management.base import BaseCommand, CommandError

from content.models import ZODIAC_SIGNS

_HEADER = ["date", "sign", "forecast", "advice"]

_INSTRUCTIONS = [
    "Инструкция для копирайтера",
    "",
    "Заполняйте лист «Гороскопы»: на каждую дату ровно по 12 строк — по одному знаку.",
    "",
    "Колонки:",
    "  date     — дата, не менять (формат ГГГГ-ММ-ДД)",
    "  sign     — знак зодиака, не менять",
    "  forecast — основной текст гороскопа, 2–4 предложения",
    "  advice   — совет в одно предложение (можно оставить пустым)",
    "",
    "Требования:",
    "  • нейтральный, приятный тон; без запугивания и категоричных обещаний",
    "  • избегать тем здоровья и «железных» финансовых прогнозов",
    "  • тексты не должны повторяться у разных знаков в один день",
    "  • не использовать чужие тексты (только оригинальные)",
    "",
    "После сдачи файла его импортируют в систему: он придёт черновиком,",
    "пройдёт вычитку и будет включён на экраны.",
]

_RULES = {
    "Астрологический прогноз не является научным фактом",
    }


def _month_days(year: int, month: int) -> int:
    """Возвращает количество дней в месяце.

    Args:
        year: Год.
        month: Месяц (1–12).

    Returns:
        Число дней в указанном месяце.
    """
    return monthrange(year, month)[1]


class Command(BaseCommand):
    help = "Генерация Excel-шаблона гороскопов для копирайтера"

    def add_arguments(self, parser):
        """Задаёт аргументы командной строки.

        Args:
            parser: Парсер аргументов команды.
        """
        parser.add_argument("--month", help="Месяц в формате ГГГГ-ММ (по умолчанию — следующий)")
        parser.add_argument("output", nargs="?", help="Путь к файлу (по умолчанию horoscope_ГГГГ-ММ.xlsx)")

    def handle(self, *args, **options):
        """Создаёт Excel-шаблон гороскопов и сохраняет его на диск."""
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
        from openpyxl.utils import get_column_letter

        if options["month"]:
            try:
                year, month = (int(x) for x in options["month"].split("-"))
                date(year, month, 1)
            except ValueError as exc:
                raise CommandError("--month должен быть в формате ГГГГ-ММ") from exc
        else:
            today = date.today()
            year, month = (today.year, today.month + 1) if today.month < 12 else (today.year + 1, 1)

        output = options["output"] or f"horoscope_{year}-{month:02d}.xlsx"
        if os.path.exists(output):
            raise CommandError(f"Файл уже существует: {output}")

        wb = Workbook()
        ws = wb.active
        ws.title = "Гороскопы"

        ws.append(_HEADER)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"

        empty = ""
        for day in range(1, _month_days(year, month) + 1):
            day_iso = date(year, month, day).isoformat()
            for _code, label in ZODIAC_SIGNS:
                ws.append([day_iso, label, empty, empty])

        widths = {1: 14, 2: 16, 3: 60, 4: 40}
        for col, width in widths.items():
            ws.column_dimensions[get_column_letter(col)].width = width
        for row in ws.iter_rows(min_row=2, min_col=1, max_col=4):
            row[2].alignment = Alignment(wrap_text=True, vertical="top")
            row[3].alignment = Alignment(wrap_text=True, vertical="top")

        instructions = wb.create_sheet("ТЗ (инструкция)")
        for line in _INSTRUCTIONS:
            instructions.append([line])
        instructions.column_dimensions["A"].width = 90

        wb.save(output)
        self.stdout.write(self.style.SUCCESS(f"Шаблон создан: {output}"))
        self.stdout.write(
            f"Строк данных: {_month_days(year, month)} дней x 12 знаков = "
            f"{_month_days(year, month) * len(ZODIAC_SIGNS)}."
        )