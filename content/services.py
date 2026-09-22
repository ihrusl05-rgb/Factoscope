"""Парсинг, валидация и запись импортируемого контента.

Поддерживаются форматы: JSON, XML, Excel (.xlsx) и CSV.
"""

from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from defusedxml import ElementTree as DefusedET

from .models import Fact, Horoscope, ImportLog, normalize_sign

REQUIRED_HOROSCOPE_FIELDS = ("sign", "date", "text")
REQUIRED_FACT_FIELDS = ("text",)


class ImportErrorValue(Exception):
    """Ошибка валидации отдельной записи."""


class ImportParseError(Exception):
    """Файл не удалось разобрать или он имеет неверную структуру."""


def _strip(value: Any) -> str:
    """Приводит значение к строке и убирает пробелы по краям.

    Args:
        value: Любое значение (в том числе ``None``).

    Returns:
        Строка; ``None`` превращается в ``""``.
    """
    return str(value or "").strip()


_HORO_TEXT_ALIASES = ("text", "forecast", "prediction", "предсказание", "прогноз", "описание")
_HORO_ADVICE_ALIASES = ("advice", "совет", "рекомендация")


def _normalize_horoscope_json(record, default_date) -> dict:
    """Приводит JSON-запись гороскопа к единому виду ``{sign, date, text}``.

    Принимает разные имена полей: ``text``/``forecast``/``prediction``
    и ``advice``/``совет`` (объединяются в ``text`` через пустую строку,
    как в табличных форматах). Дату берёт из записи, а если её нет —
    использует дату из корня файла.

    Args:
        record: Словарь с полями гороскопа.
        default_date: Дата по умолчанию из корня файла (может быть ``None``).

    Returns:
        Словарь с ключами ``sign``, ``date``, ``text``.

    Raises:
        ImportParseError: Если ``record`` — не словарь.
    """
    if not isinstance(record, dict):
        raise ImportParseError(f"Запись гороскопа должна быть объектом, получено: {record!r}")
    text = next((_strip(record[k]) for k in _HORO_TEXT_ALIASES if _strip(record.get(k))), "")
    advice = next((_strip(record[k]) for k in _HORO_ADVICE_ALIASES if _strip(record.get(k))), "")
    record = dict(record)
    record["text"] = _combine_text(text, advice) or text or advice
    if not _strip(record.get("date")) and default_date:
        record["date"] = default_date
    return record


def validate_horoscope(record: dict) -> dict:
    """Проверяет запись гороскопа и возвращает нормализованные поля.

    Args:
        record: Словарь с ключами ``sign``, ``date``, ``text``.

    Returns:
        Словарь ``{"sign", "date", "text"}`` с приведёнными значениями.

    Raises:
        ImportErrorValue: Если отсутствуют обязательные поля, знак не
            распознан или дата неверна.
    """
    missing = [f for f in REQUIRED_HOROSCOPE_FIELDS if not _strip(record.get(f, ""))]
    if missing:
        raise ImportErrorValue(f"Нет обязательных полей: {', '.join(missing)}")
    sign = normalize_sign(_strip(record["sign"]))
    if not sign:
        raise ImportErrorValue(f"Неизвестный знак: {_strip(record['sign'])!r}")
    try:
        value = record["date"]
        if isinstance(value, datetime):
            dt = value.date()
        elif isinstance(value, date):
            dt = value
        else:
            dt = date.fromisoformat(_strip(value))
    except ValueError as exc:
        raise ImportErrorValue(f"Неверная дата: {record['date']!r}") from exc
    return {"sign": sign, "date": dt, "text": _strip(record["text"])}


def validate_fact(record: dict) -> dict:
    """Проверяет запись факта и возвращает нормализованные поля.

    Args:
        record: Словарь с ключами ``text``, ``category``, ``ordering``.

    Returns:
        Словарь ``{"text", "category", "ordering"}`` с приведёнными значениями.

    Raises:
        ImportErrorValue: Если текст пуст или порядок не является числом.
    """
    if not _strip(record.get("text", "")):
        raise ImportErrorValue("Пустой текст факта")
    try:
        ordering = int(_strip(record.get("ordering", "0")) or 0)
    except ValueError as exc:
        raise ImportErrorValue(f"Неверный порядок: {record.get('ordering')!r}") from exc
    return {
        "text": _strip(record["text"]),
        "category": _strip(record.get("category", "")),
        "ordering": ordering,
    }


def parse_content(source: bytes, fmt: str) -> dict:
    """Разбирает байты файла на записи гороскопов и фактов.

    Args:
        source: Содержимое файла в байтах.
        fmt: Формат: ``json``, ``xml``, ``xlsx`` или ``csv``.

    Returns:
        Словарь ``{"horoscopes": [...], "facts": [...]}`` со «сырыми»
        записями (до валидации).

    Raises:
        ImportParseError: Если формат неизвестен, структура неверна или
            в файле нет ни одной записи.
    """
    fmt = fmt.lower()
    if fmt == "json":
        try:
            data = json.loads(source)
        except json.JSONDecodeError as exc:
            raise ImportParseError(f"Некорректный JSON: {exc}") from exc
        if not isinstance(data, dict):
            raise ImportParseError("Корень JSON должен быть объектом")
        horoscopes = data.get("horoscopes") or data.get("horoscope") or []
        facts = data.get("facts") or []
        default_date = data.get("date")
        horoscopes = [_normalize_horoscope_json(rec, default_date) for rec in horoscopes]
    elif fmt == "xml":
        horoscopes, facts = _parse_xml_block(source)
    elif fmt == "xlsx":
        horoscopes = _parse_xlsx_horoscopes(source)
        facts = []
    elif fmt == "csv":
        horoscopes = _parse_csv_horoscopes(source)
        facts = []
    else:
        raise ImportParseError(f"Неизвестный формат: {fmt}")
    if not horoscopes and not facts:
        raise ImportParseError(f"В файле формата {fmt} не найдено записей (гороскопы/факты)")
    return {"horoscopes": horoscopes, "facts": facts}


# ---------------------------------------------------------------------------
# Табличные форматы (Excel/.csv) для копирайтера
# ---------------------------------------------------------------------------

_DATE_ALIASES = {"date", "дата", "день", "число"}
_SIGN_ALIASES = {"sign", "знак", "знак зодиака", "зодиак", "name", "имя"}
_TEXT_ALIASES = {"forecast", "text", "текст", "предсказание", "прогноз", "описание"}
_ADVICE_ALIASES = {"advice", "совет", "рекомендация"}
_RUSSIAN_MONTHS = {
    "январь": 1,
    "января": 1,
    "февраль": 2,
    "февраля": 2,
    "март": 3,
    "марта": 3,
    "апрель": 4,
    "апреля": 4,
    "май": 5,
    "мая": 5,
    "июнь": 6,
    "июня": 6,
    "июль": 7,
    "июля": 7,
    "август": 8,
    "августа": 8,
    "сентябрь": 9,
    "сентября": 9,
    "октябрь": 10,
    "октября": 10,
    "ноябрь": 11,
    "ноября": 11,
    "декабрь": 12,
    "декабря": 12,
}


def _norm_col(value) -> str:
    return str(value or "").strip().lower()


def _resolve_columns(header) -> dict:
    if not header:
        return {"date": None, "sign": None, "text": None, "advice": None}
    cols = {i: _norm_col(v) for i, v in enumerate(header)}
    found = {"date": None, "sign": None, "text": None, "advice": None}
    for i, name in cols.items():
        if found["date"] is None and name in _DATE_ALIASES:
            found["date"] = i
        elif found["sign"] is None and name in _SIGN_ALIASES:
            found["sign"] = i
        elif found["text"] is None and name in _TEXT_ALIASES:
            found["text"] = i
        elif found["advice"] is None and name in _ADVICE_ALIASES:
            found["advice"] = i
    return found


def _combine_text(*parts) -> str:
    """Склеивает непустые части текста через пустую строку.

    Args:
        *parts: Значения, которые нужно объединить.

    Returns:
        Строка с непустыми частями, разделёнными пустой строкой.
    """
    vals = []
    for p in parts:
        if p is None:
            continue
        s = str(p).strip()
        if s:
            vals.append(s)
    return "\n\n".join(vals)


def _parse_header_date(value) -> date | None:
    """Извлекает дату из заголовка столбца таблицы.

    Принимает ``datetime``/``date``, формат ``ДД.ММ.ГГГГ``
    или русский вариант «16 сентября 2026».

    Args:
        value: Значение ячейки заголовка.

    Returns:
        ``date`` или ``None``, если дата не распознана.
    """
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    text = _norm_col(value)
    numeric = re.search(r"(?<!\d)(\d{1,2})\.(\d{1,2})\.(\d{4})(?!\d)", text)
    if numeric:
        try:
            return date(int(numeric.group(3)), int(numeric.group(2)), int(numeric.group(1)))
        except ValueError:
            return None

    words = re.search(r"(?<!\d)(\d{1,2})\s+([а-яё]+)\s+(\d{4})(?!\d)", text)
    if not words:
        return None
    month = _RUSSIAN_MONTHS.get(words.group(2))
    if month is None:
        return None
    try:
        return date(int(words.group(3)), month, int(words.group(1)))
    except ValueError:
        return None


def _table_to_horoscopes(rows, *, source_name: str) -> list[dict]:
    """Разбирает таблицу (xlsx/csv) в записи гороскопов.

    Поддерживается длинный формат (`date | sign | forecast | advice`)
    и широкий (`знак зодиака | период | Гороскоп на DD.MM.YYYY | ...`).

    Args:
        rows: Итератор строк таблицы.
        source_name: Имя источника (для сообщений об ошибках).

    Returns:
        Список «сырых» записей ``{"sign", "date", "text"}``.

    Raises:
        ImportParseError: Если шапка не распознана или нет строк с данными.
    """
    it = iter(rows)
    header = next(it, None)
    cols = _resolve_columns(header)

    # --- Длинный формат ---
    missing = [k for k, v in cols.items() if v is None and k != "advice"]
    if not missing:
        records = []
        for row in it:
            if not row or not any(str(v).strip() for v in row if v is not None):
                continue
            text = _combine_text(row[cols["text"]], row[cols["advice"]] if cols["advice"] is not None else None)
            if not text:
                continue
            records.append({"sign": row[cols["sign"]], "date": row[cols["date"]], "text": text})
        if records:
            return records

    # --- Широкий формат: даты в заголовках колонок ---
    date_cols: list[tuple[int, date]] = []
    if header:
        for i, h in enumerate(header):
            dt = _parse_header_date(h)
            if dt is not None:
                date_cols.append((i, dt))

    if not date_cols:
        raise ImportParseError(
            f"Шапка таблицы должна содержать колонки: date/дата, sign/знак, "
            f"forecast/текст (опц. advice/совет) или формат «широкая таблица» "
            f"с датами в заголовках (напр. «Гороскоп на 16.09.2026»)"
        )

    sign_col = cols.get("sign") if cols.get("sign") is not None else 0
    records = []
    for row in it:
        if not row or not any(str(v).strip() for v in row if v is not None):
            continue
        sign_val = row[sign_col]
        if not str(sign_val or "").strip():
            continue
        for col_i, dt in date_cols:
            if col_i >= len(row):
                continue
            text = str(row[col_i] or "").strip()
            if text:
                records.append({"sign": sign_val, "date": dt, "text": text})
    if not records:
        raise ImportParseError(f"Таблица {source_name} не содержит строк с данными")
    return records


def _parse_xlsx_horoscopes(source: bytes) -> list[dict]:
    """Разбирает .xlsx-файл в записи гороскопов.

    Args:
        source: Содержимое файла в байтах.

    Returns:
        Список «сырых» записей гороскопов.

    Raises:
        ImportParseError: Если файл не является корректным .xlsx.
    """
    from io import BytesIO

    from openpyxl import load_workbook

    try:
        wb = load_workbook(BytesIO(source), data_only=True, read_only=True)
    except Exception as exc:
        raise ImportParseError(f"Некорректный .xlsx: {exc}") from exc
    ws = wb.active
    return _table_to_horoscopes(ws.iter_rows(values_only=True), source_name=ws.title)


def _detect_delimiter(lines: list[str]) -> str:
    """Выбирает разделитель CSV по частоте встречаемости в тексте.

    Args:
        lines: Строки текста.

    Returns:
        Разделитель: ``";"``, ``","``, ``"\\t"`` или ``"|"``.
    """
    best, best_count = ";", 0
    for d in (";", ",", "\t", "|"):
        count = sum(line.count(d) for line in lines[:20])
        if count > best_count:
            best, best_count = d, count
    return best


def _parse_csv_horoscopes(source: bytes) -> list[dict]:
    """Разбирает CSV в записи гороскопов.

    Args:
        source: Содержимое файла в байтах (UTF-8, опционально с BOM).

    Returns:
        Список «сырых» записей гороскопов.

    Raises:
        ImportParseError: Если файл не в UTF-8 или CSV повреждён.
    """
    import csv
    import io as _io

    try:
        text = source.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ImportParseError(f"CSV должен быть в кодировке UTF-8: {exc}") from exc
    delimiter = _detect_delimiter(text.splitlines())
    try:
        reader = csv.reader(_io.StringIO(text), delimiter=delimiter)
        rows = [list(row) for row in reader]
    except Exception as exc:
        raise ImportParseError(f"Некорректный CSV: {exc}") from exc
    return _table_to_horoscopes(rows, source_name="CSV")


_TEXT_TAGS = ("forecast", "advice", "text", "description")


def _sign_text(element) -> str:
    """Собирает текст гороскопа из дочерних тегов.

    Args:
        element: XML-элемент с тегами ``<forecast>``/``<advice>``/``<text>``.

    Returns:
        Текст, объединённый пустой строкой; может быть ``""``.
    """
    parts = []
    for child in element:
        if child.tag in _TEXT_TAGS:
            parts.append(" ".join((child.text or "").split()))
    return "\n\n".join(p for p in parts if p)


def _parse_xml_block(source: bytes) -> tuple[list, list]:
    """Разбирает XML в пару ``(horoscopes, facts)``.

    Принимает разные структуры:

    Вариант A — плоский::

        <content>
          <horoscope sign="aries" date="2026-09-16">текст</horoscope>
          <fact category="science">текст факта</fact>
        </content>

    Вариант B — сводка от поставщика (12 знаков)::

        <horoscope date="2026-09-16" day="среда">
          <sign name="Овен">
            <forecast>...</forecast>
            <advice>...</advice>
          </sign>
        </horoscope>

    Args:
        source: Содержимое файла в байтах.

    Returns:
        Кортеж ``(horoscopes, facts)`` со «сырыми» записями.

    Raises:
        ImportParseError: Если XML некорректен или небезопасен.
    """
    try:
        root = DefusedET.fromstring(source)
    except Exception as exc:
        raise ImportParseError(f"Некорректный XML: {exc}") from exc

    horoscopes: list[dict] = []
    facts: list[dict] = []
    for el in root.iter():
        attrs = dict(el.attrib or {})
        if el.tag == "horoscope":
            date_attr = attrs.get("date")
            sign_attr = attrs.get("sign")
            signs = [c for c in el if c.tag == "sign"]
            if date_attr and signs:
                # Вариант B: контейнер с <sign name="...">
                for sign_el in signs:
                    text = _sign_text(sign_el)
                    if text:
                        horoscopes.append(
                            {"sign": sign_el.attrib.get("name", ""), "date": date_attr, "text": text}
                        )
            elif date_attr and sign_attr:
                # Вариант A: одиночный <horoscope ...>текст</horoscope>
                text = " ".join((el.text or "").split())
                if not text:
                    text = _sign_text(el)
                if text:
                    horoscopes.append({"sign": sign_attr, "date": date_attr, "text": text})
        elif el.tag == "fact":
            rec = dict(attrs)
            rec.setdefault("text", " ".join((el.text or "").split()))
            if not rec["text"]:
                rec["text"] = _sign_text(el)
            facts.append(rec)
    return horoscopes, facts


@dataclass
class ImportResult:
    added: int = 0
    updated: int = 0
    unchanged: int = 0
    errors: list = field(default_factory=list)
    records: dict = field(default_factory=dict)

    def total(self) -> int:
        """Суммарное количество обработанных записей.

        Returns:
            ``added + updated + unchanged``.
        """
        return self.added + self.updated + self.unchanged


class ContentImporter:
    """Валидирует записи, делает diff с БД и применяет изменения.

    Можно использовать и как «предпросмотр» (`preview_rows`/`preview_summary`),
    и для фактического сохранения (`commit`).
    """

    def __init__(
        self,
        source: bytes,
        fmt: str,
        *,
        force_update_text: bool = True,
        mark_draft: bool = False,
    ):
        """Создаёт импортёр из байтов файла.

        Args:
            source: Содержимое файла в байтах.
            fmt: Формат: ``json``, ``xml``, ``xlsx`` или ``csv``.
            force_update_text: Перезаписывать ли текст существующих
                гороскопов при совпадении ``(sign, date)``.
            mark_draft: Помечать ли новые гороскопы черновиком.
        """
        self.fmt = fmt.lower()
        parsed = parse_content(source, self.fmt)
        self.horoscopes: list[dict] = []
        self.facts: list[dict] = []
        self.errors: list[tuple[str, int, str]] = []
        self.force_update_text = force_update_text
        self.mark_draft = mark_draft

        for i, rec in enumerate(parsed["horoscopes"], start=1):
            try:
                self.horoscopes.append(validate_horoscope(rec))
            except ImportErrorValue as exc:
                self.errors.append(("horoscope", i, str(exc)))

        for i, rec in enumerate(parsed["facts"], start=1):
            try:
                self.facts.append(validate_fact(rec))
            except (ImportErrorValue, ValueError) as exc:
                self.errors.append(("fact", i, str(exc)))

    def has_valid_records(self) -> bool:
        """Есть ли хотя бы одна валидная запись (гороскоп или факт).

        Returns:
            ``True``, если файл можно импортировать.
        """
        return bool(self.horoscopes or self.facts)

    def _diff_rows(self) -> list[dict]:
        """Строит полный список строк предпросмотра.

        Returns:
            Список словарей со статусами ``add``/``update``/``unchanged``/``error``.
        """
        rows = [
            {"kind": kind, "row": i, "status": "error", "detail": err}
            for kind, i, err in self.errors
        ]
        for rec in self.horoscopes:
            rows.append(self._diff_horoscope(rec))
        for rec in self.facts:
            rows.append(self._diff_fact(rec))
        return rows

    def _diff_horoscope(self, rec: dict) -> dict:
        """Определяет статус записи гороскопа относительно БД.

        Args:
            rec: Валидированная запись ``{"sign", "date", "text"}``.

        Returns:
            Словарь ``{"kind", "key", "status"}``.
        """
        exists = Horoscope.objects.filter(sign=rec["sign"], date=rec["date"]).first()
        key = f"{rec['sign']} · {rec['date']}"
        if exists is None:
            return {"kind": "horoscope", "key": key, "status": "add"}
        if self.force_update_text and exists.text != rec["text"]:
            return {"kind": "horoscope", "key": key, "status": "update"}
        return {"kind": "horoscope", "key": key, "status": "unchanged"}

    def _diff_fact(self, rec: dict) -> dict:
        """Определяет статус записи факта относительно БД.

        Args:
            rec: Валидированная запись ``{"text", "category", "ordering"}``.

        Returns:
            Словарь ``{"kind", "key", "status"}``.
        """
        if Fact.objects.filter(text=rec["text"]).exists():
            return {"kind": "fact", "key": rec["text"][:45], "status": "unchanged"}
        return {"kind": "fact", "key": rec["text"][:45], "status": "add"}

    def preview_rows(self) -> list[dict]:
        """Возвращает построчную разницу файла с базой.

        Returns:
            Список строк предпросмотра (см. ``_diff_rows``).
        """
        return self._diff_rows()

    def preview_summary(self) -> dict:
        """Возвращает сводку по статусам всех записей.

        Returns:
            Словарь со счётчиками ``add``/``update``/``unchanged``/``error``.
        """
        rows = self._diff_rows()
        return {status: sum(1 for r in rows if r["status"] == status) for status in ("add", "update", "unchanged", "error")}

    def commit(self) -> ImportResult:
        """Применяет изменения в базу данных.

        Returns:
            Объект :class:`ImportResult` со счётчиками ``added``/``updated``/
            ``unchanged`` и списком ``errors``.
        """
        result = ImportResult()
        result.errors = [f"{kind} #{i}: {err}" for kind, i, err in self.errors]

        for rec in self.horoscopes:
            diff = self._diff_horoscope(rec)
            if diff["status"] == "add":
                Horoscope.objects.create(
                    sign=rec["sign"],
                    date=rec["date"],
                    text=rec["text"],
                    is_draft=self.mark_draft,
                )
                result.added += 1
            elif diff["status"] == "update":
                Horoscope.objects.filter(sign=rec["sign"], date=rec["date"]).update(text=rec["text"])
                result.updated += 1
            else:
                result.unchanged += 1

        for rec in self.facts:
            if Fact.objects.filter(text=rec["text"]).exists():
                result.unchanged += 1
            else:
                Fact.objects.create(text=rec["text"], category=rec["category"], ordering=rec["ordering"])
                result.added += 1

        result.records = {"horoscopes": len(self.horoscopes), "facts": len(self.facts)}
        return result

    @staticmethod
    def save_log(filename: str, fmt: str, result: ImportResult) -> ImportLog:
        """Сохраняет запись об импорте в журнал.

        Args:
            filename: Имя файла-источника.
            fmt: Формат файла (``json``, ``xml``, ``xlsx``, ``csv``).
            result: Результат импорта.

        Returns:
            Созданная запись :class:`ImportLog`.
        """
        status = ImportLog.Status.ERROR
        if result.added or result.updated:
            status = ImportLog.Status.SUCCESS if not result.errors else ImportLog.Status.PARTIAL
        return ImportLog.objects.create(
            filename=filename,
            format=fmt.lower(),
            status=status,
            report={
                "added": result.added,
                "updated": result.updated,
                "unchanged": result.unchanged,
                "errors": result.errors,
                "records": result.records,
            },
        )
